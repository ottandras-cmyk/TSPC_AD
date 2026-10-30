import os, json, glob, subprocess, time, urllib.request, urllib.error, socket, ssl, threading, re, logging, hashlib
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import JSONResponse, Response, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Text, cast, desc, or_, text
from sqlalchemy.orm import Session
from datetime import datetime
from io import BytesIO
from database import get_db
import models
from services.sharepoint_service import (
    SharePointIntegrationError,
    SHAREPOINT_LIST_MODELS,
    list_sharepoint_lists,
    sync_inventory_to_database,
    sync_sharepoint_list_to_database,
    sync_sharepoint_lists_to_database,
    test_sharepoint_connection,
)

router = APIRouter(prefix="/api/settings", tags=["settings"])
logger = logging.getLogger(__name__)

class SharePointImageSyncRequest(BaseModel):
    username: str = Field(min_length=1, max_length=256)
    password: str = Field(min_length=1, max_length=1024)


def get_all_settings():
    raw_data = {}
    for sf in ["app/config/settings.json", "config/settings.json", "app/settings.json", "settings.json"]:
        if os.path.exists(sf):
            try:
                with open(sf, "r", encoding="utf-8", errors="ignore") as f:
                    content = json.load(f)
                    if isinstance(content, dict) and len(content) > 0:
                        raw_data.update(content)
                        break
            except Exception: pass

    for ef in [".env", "app/.env"]:
        if os.path.exists(ef):
            try:
                with open(ef, "r", encoding="utf-8", errors="ignore") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            raw_data[k.strip()] = v.strip().strip('"').strip("'")
            except Exception: pass

    norm = {
        "status": "SUCCESS",
        "easyproject": {
            "url": "https://ep.tspc.hu/",
            "api_key": "",
            "project_id": "2403",
            "sync_interval_minutes": 15
        },
        "ad": {
            "server": "dc1.tspc.hu:389",
            "base_dn": "DC=tspc,DC=hu",
            "bind_user": "",
            "bind_password": "",
            "sync_interval_minutes": 15
        },
        "telegram": {
            "bot_token": "",
            "chat_id": "",
            "notify_helpdesk": True,
            "notify_audit": True
        },
        "github": {
            "repo_url": "https://github.com/tspcit/ad-tracker.git",
            "token": "",
            "branch": "main",
            "auto_tag": True
        },
        "sharepoint": {
            "tenant_id": "",
            "client_id": "",
            "client_secret": "",
            "site_url": "https://tspckft.sharepoint.com/sites/TSPCIT",
            "sync_enabled": False,
            "sync_interval_minutes": 60,
        }
    }

    for sec in ["easyproject", "ad", "telegram", "github"]:
        if sec in raw_data and isinstance(raw_data[sec], dict):
            norm[sec].update(raw_data[sec])
    if isinstance(raw_data.get("sharepoint"), dict):
        norm["sharepoint"].update(raw_data["sharepoint"])

    def scan_dict(d):
        for k, v in d.items():
            if isinstance(v, dict):
                scan_dict(v)
            else:
                k_u = k.upper()
                val = str(v).strip() if v is not None else ""
                if not val or val.startswith("•••"): continue
                
                # Telegram
                if any(x in k_u for x in ["TG_BOT_TOKEN", "TELEGRAM_BOT_TOKEN", "BOT_TOKEN", "TG_TOKEN"]) and not norm["telegram"]["bot_token"]:
                    norm["telegram"]["bot_token"] = val
                elif any(x in k_u for x in ["TG_CHAT_ID", "TELEGRAM_CHAT_ID", "CHAT_ID", "TG_CHAT"]) and not norm["telegram"]["chat_id"]:
                    norm["telegram"]["chat_id"] = val
                elif "TG_EV_SYNC_ERROR" in k_u:
                    norm["telegram"]["notify_helpdesk"] = bool(v)
                elif "NOTIFY_AUDIT" in k_u:
                    norm["telegram"]["notify_audit"] = bool(v)
                    
                # EasyProject
                elif any(x in k_u for x in ["EP_URL", "EASYPROJECT_URL", "EASY_PROJECT_URL"]):
                    norm["easyproject"]["url"] = val
                elif any(x in k_u for x in ["EP_API_KEY", "EP_KEY", "EASYPROJECT_API_KEY", "EP_TOKEN"]) and not norm["easyproject"]["api_key"]:
                    norm["easyproject"]["api_key"] = val
                elif any(x in k_u for x in ["EP_PROJECT_ID", "PROJECT_ID"]):
                    norm["easyproject"]["project_id"] = val
                elif "EP_SYNC_INTERVAL" in k_u or "EP_INTERVAL" in k_u:
                    try: norm["easyproject"]["sync_interval_minutes"] = int(v)
                    except: pass
                    
                # AD
                elif any(x in k_u for x in ["AD_SERVER", "LDAP_SERVER", "AD_HOST"]):
                    norm["ad"]["server"] = val
                elif any(x in k_u for x in ["AD_BASE_DN", "LDAP_BASE_DN", "BASE_DN"]):
                    norm["ad"]["base_dn"] = val
                elif any(x in k_u for x in ["AD_BIND_USER", "LDAP_BIND_USER", "AD_USER", "BIND_USER"]):
                    norm["ad"]["bind_user"] = val
                elif any(x in k_u for x in ["AD_BIND_PASSWORD", "LDAP_BIND_PASSWORD", "AD_PASSWORD", "BIND_PASSWORD"]):
                    norm["ad"]["bind_password"] = val
                elif "AD_SYNC_INTERVAL" in k_u or "AD_INTERVAL" in k_u or "SYNC_INTERVAL_MINUTES" in k_u:
                    try: norm["ad"]["sync_interval_minutes"] = int(v)
                    except: pass
                    
                # GitHub
                elif any(x in k_u for x in ["GITHUB_TOKEN", "GH_TOKEN", "PAT"]) and not norm["github"]["token"]:
                    norm["github"]["token"] = val
                elif any(x in k_u for x in ["GITHUB_REPO", "REPO_URL"]):
                    norm["github"]["repo_url"] = val

    scan_dict(raw_data)
    norm["GITHUB_TOKEN"] = norm["github"]["token"]
    norm["tab_visibility"] = raw_data.get("tab_visibility", {})
    return norm

@router.get("")
@router.get("/")
def get_settings():
    settings = get_all_settings()
    sharepoint = dict(settings["sharepoint"])
    sharepoint["secret_configured"] = bool(sharepoint.get("client_secret"))
    sharepoint["client_secret"] = ""
    settings["sharepoint"] = sharepoint
    return settings


def _require_settings_user(request: Request) -> None:
    from auth import COOKIE_NAME, verify_session_token

    user, _ = verify_session_token(request.cookies.get(COOKIE_NAME))
    if not user:
        raise HTTPException(status_code=401, detail="Bejelentkezés szükséges.")


def _require_secure_transport(request: Request) -> None:
    forwarded_proto = request.headers.get("x-forwarded-proto", "").split(",", 1)[0].strip().lower()
    if request.url.scheme != "https" and forwarded_proto != "https":
        raise HTTPException(
            status_code=400,
            detail="A SharePoint-képszinkron hitelesítési adatokat kér, ezért csak HTTPS-kapcsolaton indítható.",
        )


def _save_sharepoint_config(config: dict) -> bool:
    settings_files = ["app/config/settings.json", "config/settings.json", "app/settings.json", "settings.json"]
    existing = {}
    for settings_file in settings_files:
        if os.path.exists(settings_file):
            try:
                with open(settings_file, "r", encoding="utf-8") as source:
                    loaded = json.load(source)
                if isinstance(loaded, dict):
                    existing = loaded
                    break
            except Exception:
                continue

    existing["sharepoint"] = config
    saved = False
    for settings_file in settings_files:
        try:
            directory = os.path.dirname(settings_file)
            if directory:
                os.makedirs(directory, exist_ok=True)
            with open(settings_file, "w", encoding="utf-8") as destination:
                json.dump(existing, destination, indent=2, ensure_ascii=False)
            saved = True
        except OSError:
            continue
    return saved


def _sharepoint_request_config(body: dict) -> dict:
    config = dict(get_all_settings()["sharepoint"])
    for key in ("tenant_id", "client_id", "site_url"):
        if key in body:
            config[key] = str(body[key] or "").strip()
    if body.get("client_secret"):
        config["client_secret"] = str(body["client_secret"]).strip()
    if "sync_enabled" in body:
        if not isinstance(body["sync_enabled"], bool):
            raise ValueError("Az automatikus SharePoint-frissítés kapcsolójának logikai értéknek kell lennie.")
        config["sync_enabled"] = body["sync_enabled"]
    if "sync_interval_minutes" in body:
        try:
            interval = int(body["sync_interval_minutes"])
        except (TypeError, ValueError) as exc:
            raise ValueError("Érvénytelen SharePoint-frissítési gyakoriság.") from exc
        if interval not in {5, 15, 30, 60, 180, 360, 720, 1440}:
            raise ValueError("A SharePoint-frissítés gyakorisága nem támogatott.")
        config["sync_interval_minutes"] = interval
    return config


@router.get("/sharepoint/catalog")
def get_sharepoint_catalog(request: Request):
    _require_settings_user(request)
    try:
        return list_sharepoint_lists(get_all_settings()["sharepoint"])
    except SharePointIntegrationError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})


@router.get("/sharepoint/data/{list_name}")
def get_sharepoint_list_data(
    request: Request,
    list_name: str,
    search: str = "",
    sort_by: str = "sharepoint_item_id",
    sort_direction: str = "asc",
    page: int = 1,
    page_size: int = 50,
    show_deleted: bool = False,
    db: Session = Depends(get_db),
):
    _require_settings_user(request)
    model = SHAREPOINT_LIST_MODELS.get(list_name)
    if model is None:
        return JSONResponse(status_code=404, content={"status": "ERROR", "message": "Ismeretlen SharePoint-lista."})

    page = max(page, 1)
    page_size = min(max(page_size, 1), 200)
    query = db.query(model)
    if not show_deleted:
        query = query.filter(model.is_deleted.is_(False))
    if search.strip():
        search_value = f"%{search.strip()}%"
        query = query.filter(
            or_(
                model.sharepoint_item_id.ilike(search_value),
                cast(model.fields, Text).ilike(search_value),
            )
        )

    total = query.count()
    direction = sort_direction.casefold()
    if sort_by == "sharepoint_item_id":
        sort_column = model.sharepoint_item_id
    elif sort_by == "source_modified_at":
        sort_column = model.source_modified_at
    elif sort_by == "synced_at":
        sort_column = model.synced_at
    elif sort_by == "is_deleted":
        sort_column = model.is_deleted
    else:
        sort_column = model.fields[sort_by].astext
    ordering = sort_column.desc() if direction == "desc" else sort_column.asc()
    records = query.order_by(ordering.nullslast(), model.sharepoint_item_id.asc()).offset(
        (page - 1) * page_size
    ).limit(page_size).all()
    field_rows = db.execute(text(
        f'SELECT DISTINCT keys.field_name FROM "{model.__tablename__}" '
        "CROSS JOIN LATERAL jsonb_object_keys(fields) AS keys(field_name) "
        "ORDER BY keys.field_name"
    )).all()

    return {
        "status": "SUCCESS",
        "list_name": list_name,
        "field_names": [row[0] for row in field_rows],
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [
            {
                "sharepoint_item_id": record.sharepoint_item_id,
                "fields": record.fields or {},
                "source_modified_at": record.source_modified_at,
                "synced_at": record.synced_at,
                "is_deleted": record.is_deleted,
            }
            for record in records
        ],
    }


@router.get("/sharepoint/changes")
def get_sharepoint_changes(
    request: Request,
    search: str = "",
    list_name: str = "",
    item_id: str = "",
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
):
    _require_settings_user(request)
    page = max(page, 1)
    page_size = min(max(page_size, 1), 200)
    query = db.query(models.SharePointItemChange)
    if list_name:
        if list_name not in SHAREPOINT_LIST_MODELS:
            return JSONResponse(status_code=404, content={"status": "ERROR", "message": "Ismeretlen SharePoint-lista."})
        query = query.filter(models.SharePointItemChange.list_name == list_name)
    if item_id:
        if not list_name:
            return JSONResponse(status_code=400, content={"status": "ERROR", "message": "Az elem előzményeihez a lista neve is szükséges."})
        query = query.filter(models.SharePointItemChange.sharepoint_item_id == item_id)
    if search.strip():
        search_value = f"%{search.strip()}%"
        query = query.filter(
            or_(
                models.SharePointItemChange.list_name.ilike(search_value),
                models.SharePointItemChange.sharepoint_item_id.ilike(search_value),
                models.SharePointItemChange.change_type.ilike(search_value),
                cast(models.SharePointItemChange.changed_fields, Text).ilike(search_value),
                cast(models.SharePointItemChange.previous_fields, Text).ilike(search_value),
                cast(models.SharePointItemChange.current_fields, Text).ilike(search_value),
            )
        )

    total = query.count()
    ordering = (
        models.SharePointItemChange.changed_at.asc()
        if item_id
        else desc(models.SharePointItemChange.changed_at)
    )
    records = query.order_by(ordering, models.SharePointItemChange.id.asc()).offset(
        (page - 1) * page_size
    ).limit(page_size).all()
    return {
        "status": "SUCCESS",
        "total": total,
        "page": page,
        "page_size": page_size,
        "changes": [
            {
                "list_name": record.list_name,
                "sharepoint_item_id": record.sharepoint_item_id,
                "change_type": record.change_type,
                "changed_fields": record.changed_fields,
                "previous_fields": record.previous_fields,
                "current_fields": record.current_fields,
                "changed_at": record.changed_at,
                "trigger_type": record.trigger_type,
            }
            for record in records
        ],
    }


@router.post("/sharepoint/images/sync")
def sync_sharepoint_equipment_images(
    body: SharePointImageSyncRequest,
    request: Request,
):
    _require_settings_user(request)
    _require_secure_transport(request)
    if not body.username.strip():
        raise HTTPException(status_code=400, detail="Felhasználónév megadása szükséges.")
    from services.sharepoint_image_service import (
        ImageAuthenticationError,
        ImageSyncError,
        sync_sharepoint_images,
    )

    try:
        return sync_sharepoint_images(body.username.strip(), body.password)
    except ImageAuthenticationError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    except ImageSyncError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@router.get("/sharepoint/images/{item_id}")
def get_sharepoint_equipment_image_versions(
    item_id: str,
    request: Request,
    db: Session = Depends(get_db),
):
    _require_settings_user(request)
    if len(item_id) > 100:
        raise HTTPException(status_code=400, detail="Érvénytelen SharePoint ID.")

    from services.sharepoint_image_service import _extract_image_urls

    inventory_item = db.query(models.SP_eszkozleltar_nonIT).filter_by(
        sharepoint_item_id=item_id
    ).first()
    current_urls, _ = _extract_image_urls((inventory_item.fields or {}).get("field_2")) if inventory_item else ([], 0)
    current_hashes = {
        hashlib.sha256(url.encode("utf-8")).hexdigest()
        for url in current_urls
    }
    versions = (
        db.query(models.SharePointImageVersion)
        .filter_by(sharepoint_item_id=item_id)
        .order_by(
            models.SharePointImageVersion.source_url_sha256,
            desc(models.SharePointImageVersion.version_number),
        )
        .all()
    )
    grouped = {}
    for version in versions:
        group = grouped.setdefault(
            version.source_url_sha256,
            {
                "source_is_current": version.source_url_sha256 in current_hashes,
                "versions": [],
            },
        )
        group["versions"].append({
            "id": version.id,
            "version_number": version.version_number,
            "file_name": version.file_name,
            "content_type": version.content_type,
            "file_size_bytes": version.file_size_bytes,
            "content_sha256": version.content_sha256,
            "downloaded_at": version.downloaded_at,
            "last_verified_at": version.last_verified_at,
            "is_latest": not group["versions"],
            "image_url": f"/api/settings/sharepoint/images/version/{version.id}",
        })
    return {
        "status": "SUCCESS",
        "item_id": item_id,
        "image_groups": list(grouped.values()),
    }


@router.get("/sharepoint/images/version/{version_id}")
def get_sharepoint_equipment_image(
    version_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    _require_settings_user(request)
    version = db.query(models.SharePointImageVersion).filter_by(id=version_id).first()
    if version is None:
        raise HTTPException(status_code=404, detail="A kért képfájl nem található.")
    return Response(
        content=version.image_data,
        media_type=version.content_type,
        headers={
            "Cache-Control": "private, no-store",
            "Content-Disposition": f'inline; filename="{version.file_name}"',
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.get("/sharepoint/export/{list_name}")
def export_sharepoint_list(
    request: Request,
    list_name: str,
    show_deleted: bool = False,
    db: Session = Depends(get_db),
):
    _require_settings_user(request)
    model = SHAREPOINT_LIST_MODELS.get(list_name)
    if model is None:
        return JSONResponse(status_code=404, content={"status": "ERROR", "message": "Ismeretlen SharePoint-lista."})

    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    fields = db.execute(text(
        f'SELECT DISTINCT keys.field_name FROM "{model.__tablename__}" '
        "CROSS JOIN LATERAL jsonb_object_keys(fields) AS keys(field_name) "
        "ORDER BY keys.field_name"
    )).scalars().all()
    query = db.query(model)
    if not show_deleted:
        query = query.filter(model.is_deleted.is_(False))
    records = query.order_by(model.sharepoint_item_id.asc()).yield_per(500)

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=list_name[:31])
    worksheet.append([
        "SharePoint ID",
        "SharePoint módosítva",
        "Helyi frissítve",
        "Állapot",
        *fields,
    ])

    def excel_value(value):
        if value is None:
            return None
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        elif isinstance(value, datetime):
            value = value.isoformat()
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
            if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
                value = "'" + value
        return value

    for record in records:
        item_fields = record.fields or {}
        worksheet.append([
            excel_value(record.sharepoint_item_id),
            excel_value(record.source_modified_at),
            excel_value(record.synced_at),
            "Törölt" if record.is_deleted else "Aktív",
            *(excel_value(item_fields.get(field)) for field in fields),
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = re.sub(r"[^A-Za-z0-9_-]", "_", list_name) + ".xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/sharepoint/save")
async def save_sharepoint_settings(request: Request):
    _require_settings_user(request)
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        config = _sharepoint_request_config(body if isinstance(body, dict) else {})
    except ValueError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})
    if not _save_sharepoint_config(config):
        return JSONResponse(status_code=500, content={"status": "ERROR", "message": "A SharePoint-beállításokat nem sikerült elmenteni."})
    from scheduler.registry import configure_sharepoint_sync_job
    configure_sharepoint_sync_job(config.get("sync_enabled", False), config.get("sync_interval_minutes", 60))
    return {
        "status": "SUCCESS",
        "message": "SharePoint-beállítások elmentve.",
        "secret_configured": bool(config.get("client_secret")),
        "sync_enabled": config.get("sync_enabled", False),
        "sync_interval_minutes": config.get("sync_interval_minutes", 60),
    }


@router.post("/sharepoint/test")
async def test_sharepoint_settings(request: Request):
    _require_settings_user(request)
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        config = _sharepoint_request_config(body if isinstance(body, dict) else {})
    except ValueError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})
    try:
        return test_sharepoint_connection(config)
    except SharePointIntegrationError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})


@router.post("/sharepoint/download-inventory")
def download_sharepoint_inventory(request: Request):
    _require_settings_user(request)
    config = get_all_settings()["sharepoint"]
    try:
        return sync_inventory_to_database(config)
    except SharePointIntegrationError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})
    except Exception:
        logger.exception("SharePoint inventory synchronization failed")
        return JSONResponse(status_code=500, content={"status": "ERROR", "message": "Adatbázis-hiba az import közben."})


@router.post("/sharepoint/download-all")
def download_all_sharepoint_lists(request: Request):
    _require_settings_user(request)
    try:
        return sync_sharepoint_lists_to_database(get_all_settings()["sharepoint"])
    except SharePointIntegrationError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})
    except Exception:
        logger.exception("SharePoint list synchronization failed")
        return JSONResponse(
            status_code=500,
            content={"status": "ERROR", "message": "A négy lista szinkronizálása nem sikerült."},
        )


@router.post("/sharepoint/download-list/{list_name}")
def download_sharepoint_list(request: Request, list_name: str):
    _require_settings_user(request)
    config = get_all_settings()["sharepoint"]
    try:
        return sync_sharepoint_list_to_database(config, list_name)
    except SharePointIntegrationError as exc:
        return JSONResponse(status_code=400, content={"status": "ERROR", "message": str(exc)})
    except Exception:
        logger.exception("SharePoint list synchronization failed")
        return JSONResponse(status_code=500, content={"status": "ERROR", "message": "Adatbázis-hiba az import közben."})

@router.post("/save")
async def save_settings(req: Request):
    try: body = await req.json()
    except Exception: body = {}

    current = get_all_settings()
    for sec in ["easyproject", "ad", "telegram", "github"]:
        if sec in body and isinstance(body[sec], dict):
            current[sec].update(body[sec])

    cfg_file = "app/config/settings.json" if os.path.exists("app/config/settings.json") else "app/settings.json"
    existing = {}
    if os.path.exists(cfg_file):
        try:
            with open(cfg_file, "r", encoding="utf-8") as f:
                existing = json.load(f)
        except Exception: pass

    existing["EP_URL"] = current["easyproject"]["url"]
    existing["EP_API_KEY"] = current["easyproject"]["api_key"]
    existing["EP_PROJECT_ID"] = current["easyproject"]["project_id"]
    existing["EP_SYNC_INTERVAL_MINUTES"] = current["easyproject"]["sync_interval_minutes"]

    existing["AD_SERVER"] = current["ad"]["server"]
    existing["AD_BASE_DN"] = current["ad"]["base_dn"]
    existing["AD_BIND_USER"] = current["ad"]["bind_user"]
    existing["AD_BIND_PASSWORD"] = current["ad"]["bind_password"]
    existing["AD_SYNC_INTERVAL_MINUTES"] = current["ad"]["sync_interval_minutes"]

    existing["TG_BOT_TOKEN"] = current["telegram"]["bot_token"]
    existing["TG_CHAT_ID"] = current["telegram"]["chat_id"]
    existing["GITHUB_TOKEN"] = current["github"]["token"]
    
    existing["easyproject"] = current["easyproject"]
    existing["ad"] = current["ad"]
    existing["telegram"] = current["telegram"]
    existing["github"] = current["github"]
    if "tab_visibility" in body:
        existing["tab_visibility"] = body["tab_visibility"]
        current["tab_visibility"] = body["tab_visibility"]

    for sf in ["app/config/settings.json", "config/settings.json", "app/settings.json", "settings.json"]:
        try:
            dname = os.path.dirname(sf)
            if dname: os.makedirs(dname, exist_ok=True)
            with open(sf, "w", encoding="utf-8") as f:
                json.dump(existing, f, indent=2, ensure_ascii=False)
        except Exception: pass

    try:
        env_lines = []
        if os.path.exists(".env"):
            with open(".env", "r", encoding="utf-8") as ef:
                env_lines = [l for l in ef.read().splitlines() if not l.startswith("GITHUB_TOKEN")]
        if current["github"]["token"]:
            env_lines.append(f"GITHUB_TOKEN={current['github']['token']}")
        with open(".env", "w", encoding="utf-8") as ef:
            ef.write("\n".join(env_lines) + "\n")
    except Exception: pass

    response_settings = dict(current)
    sharepoint = dict(response_settings.get("sharepoint", {}))
    sharepoint["secret_configured"] = bool(sharepoint.get("client_secret"))
    sharepoint["client_secret"] = ""
    response_settings["sharepoint"] = sharepoint
    return JSONResponse(status_code=200, content={"status": "SUCCESS", "message": "Beállítások és paraméterek elmentve!", "settings": response_settings})

@router.post("/test-ep")
async def test_ep(req: Request):
    try: body = await req.json()
    except Exception: body = {}
    url = (body.get("url") or "").strip().rstrip("/")
    api_key = (body.get("api_key") or "").strip()
    if not url or not api_key:
        curr = get_all_settings()["easyproject"]
        url = url or curr.get("url", "")
        api_key = api_key or curr.get("api_key", "")
    if not url or not api_key:
        return {"status": "ERROR", "message": "Hiányzó URL vagy API kulcs!"}

    try:
        rq = urllib.request.Request(f"{url}/users/current.json", headers={"X-Redmine-API-Key": api_key, "User-Agent": "ADTracker/1.0"})
        ctx = ssl._create_unverified_context()
        with urllib.request.urlopen(rq, timeout=6, context=ctx) as resp:
            if resp.status == 200:
                d = json.loads(resp.read().decode())
                u = d.get("user", {})
                name = f"{u.get('firstname', '')} {u.get('lastname', '')}".strip() or u.get("login", "Ismeretlen")
                return {"status": "SUCCESS", "message": f"Sikeres kapcsolat! Bejelentkezve: {name} (ID: #{u.get('id', '-')})"}
    except Exception as e:
        return {"status": "ERROR", "message": f"Kapcsolódási hiba: {str(e)}"}

@router.post("/test-ad")
async def test_ad(req: Request):
    try: body = await req.json()
    except Exception: body = {}
    server = (body.get("server") or "").strip()
    user = (body.get("bind_user") or "").strip()
    pwd = (body.get("bind_password") or "").strip()
    if not server:
        curr = get_all_settings()["ad"]
        server = server or curr.get("server", "")
        user = user or curr.get("bind_user", "")
        pwd = pwd or curr.get("bind_password", "")
    if not server:
        return {"status": "ERROR", "message": "Nincs megadva AD szerver cím!"}

    host = server.replace("ldap://", "").replace("ldaps://", "").split(":")[0]
    port = 636 if "ldaps://" in server else (int(server.split(":")[1]) if ":" in server else 389)
    try:
        s = socket.create_connection((host, port), timeout=4)
        s.close()
    except Exception as e:
        return {"status": "ERROR", "message": f"AD szerver port ({host}:{port}) nem érhető el: {str(e)}"}

    try:
        from ldap3 import Server, Connection, ALL
        srv = Server(server, get_info=ALL, connect_timeout=4)
        conn = Connection(srv, user=user, password=pwd, auto_bind=True, timeout=4)
        conn.unbind()
        return {"status": "SUCCESS", "message": f"Sikeres AD/LDAP kapcsolat és hitelesítés ({host}:{port})!"}
    except Exception as le:
        return {"status": "SUCCESS", "message": f"AD szerver elérhető ({host}:{port})! (LDAP bind: {str(le)})"}

@router.post("/test-telegram")
async def test_telegram(req: Request):
    try: body = await req.json()
    except Exception: body = {}
    token = (body.get("bot_token") or "").strip()
    chat_id = (body.get("chat_id") or "").strip()
    if not token or not chat_id:
        curr = get_all_settings()["telegram"]
        token = token or curr.get("bot_token", "")
        chat_id = chat_id or curr.get("chat_id", "")
    if not token or not chat_id:
        return {"status": "ERROR", "message": "Hiányzó Bot Token vagy Chat ID!"}

    try:
        url = f"https://api.telegram.org/bot{token}/sendMessage"
        ts = time.strftime("%Y-%m-%d %H:%M:%S")
        payload = {
            "chat_id": chat_id,
            "text": f"🔔 <b>AD Tracker Rendszerközpont Teszt</b>\n\nIdőpont: <code>{ts}</code>\nA Telegram értesítő központ megfelelően működik!",
            "parse_mode": "HTML"
        }
        rq = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers={"Content-Type": "application/json"})
        ctx = ssl._create_unverified_context()
        with urllib.request.urlopen(rq, timeout=6, context=ctx) as resp:
            if resp.status == 200:
                return {"status": "SUCCESS", "message": "Teszt üzenet sikeresen elküldve a Telegram csoportba!"}
            return {"status": "ERROR", "message": f"Telegram API válaszkód: {resp.status}"}
    except Exception as e:
        return {"status": "ERROR", "message": f"Hiba a Telegram küldéskor: {str(e)}"}

@router.post("/test-git")
async def test_git(req: Request):
    try: body = await req.json()
    except Exception: body = {}

    cfg = get_all_settings()
    repo_url = (body.get("repo_url") or cfg["github"].get("repo_url") or "https://github.com/tspcit/ad-tracker.git").strip()
    token = (body.get("token") or cfg["github"].get("token") or "").strip()

    tag = subprocess.getoutput("git describe --tags --abbrev=0 2>/dev/null || cat app/git_version.txt 2>/dev/null || echo 'v1.10'").strip()
    log = subprocess.getoutput('git log -n 1 --pretty=format:"%h|%cd|%s" --date=format:"%Y-%m-%d %H:%M" 2>/dev/null').strip()
    parts = log.split("|") if log else ["-", "-", "-"]

    info = {
        "status": "SUCCESS",
        "version": tag or "v1.10",
        "name": parts[2] if len(parts) > 2 else "Kiadás",
        "date": parts[1] if len(parts) > 1 else "-",
        "commit": parts[0] if parts[0] else "-",
        "source": "Helyi Git tároló"
    }

    if repo_url and "github.com" in repo_url:
        m = re.search(r'github\.com[/:]([^/]+)/([^/\.]+)', repo_url)
        if m:
            owner, repo = m.group(1), m.group(2)
            headers = {"User-Agent": "ADTracker-App", "Accept": "application/vnd.github+json"}
            if token and len(token) > 10 and not token.startswith("•••"):
                headers["Authorization"] = f"token {token}"
            
            ctx = ssl._create_unverified_context()

            try:
                rq = urllib.request.Request(f"https://api.github.com/repos/{owner}/{repo}/releases/latest", headers=headers)
                with urllib.request.urlopen(rq, timeout=5, context=ctx) as r:
                    rel = json.loads(r.read().decode())
                    if rel.get("tag_name"):
                        info["version"] = rel.get("tag_name")
                        info["name"] = rel.get("name") or rel.get("tag_name")
                        info["date"] = (rel.get("published_at") or "")[:16].replace("T", " ")
                        info["source"] = f"GitHub Releases ({owner}/{repo})"
            except Exception: pass

            if info["source"] == "Helyi Git tároló":
                try:
                    rq = urllib.request.Request(f"https://api.github.com/repos/{owner}/{repo}/tags?per_page=1", headers=headers)
                    with urllib.request.urlopen(rq, timeout=5, context=ctx) as r:
                        tags = json.loads(r.read().decode())
                        if tags and isinstance(tags, list):
                            info["version"] = tags[0].get("name", info["version"])
                            info["commit"] = (tags[0].get("commit", {}).get("sha") or "")[:7] or info["commit"]
                            info["source"] = f"GitHub Tags ({owner}/{repo})"
                except Exception: pass

            try:
                rq = urllib.request.Request(f"https://api.github.com/repos/{owner}/{repo}/commits?per_page=1", headers=headers)
                with urllib.request.urlopen(rq, timeout=5, context=ctx) as r:
                    commits = json.loads(r.read().decode())
                    if commits and isinstance(commits, list):
                        c0 = commits[0]
                        info["commit"] = (c0.get("sha") or "")[:7] or info["commit"]
                        c_msg = c0.get("commit", {}).get("message", "").split("\n")[0]
                        if info["name"] in ["Kiadás", "-", ""]: info["name"] = c_msg
                        c_date = (c0.get("commit", {}).get("committer", {}).get("date") or "")[:16].replace("T", " ")
                        if info["date"] in ["-", ""]: info["date"] = c_date
                        if "Releases" not in info["source"]:
                            info["source"] = f"GitHub API ({owner}/{repo})"
            except Exception: pass

    info["message"] = f"GitHub kapcsolat sikeres! Verzió: {info['version']} ({info['name']}) | Dátum: {info['date']}"
    return JSONResponse(status_code=200, content=info)

def get_db_counts():
    counts = {}
    try:
        env = os.environ.copy()
        env["PGPASSWORD"] = "ChangeThisStrongPassword123!"
        cmd = ["psql", "-h", "db", "-U", "adaudit", "-d", "adaudit", "-t", "-A", "-c", "SELECT table_name FROM information_schema.tables WHERE table_schema='public';"]
        p = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=3)
        tables = [t.strip() for t in p.stdout.splitlines() if t.strip()]
        for t in tables:
            if re.match(r'^[a-zA-Z0-9_]+$', t):
                c_cmd = ["psql", "-h", "db", "-U", "adaudit", "-d", "adaudit", "-t", "-A", "-c", f"SELECT count(*) FROM {t};"]
                c_p = subprocess.run(c_cmd, env=env, capture_output=True, text=True, timeout=2)
                val = c_p.stdout.strip()
                if val.isdigit(): counts[t] = int(val)
    except Exception: pass
    return counts

@router.post("/sync-ad-now")
def sync_ad_now():
    t0 = time.time()
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    cmd = "python3 /app/sync_engine.py 2>&1 || python3 sync_engine.py 2>&1"
    try:
        p = subprocess.run(cmd, shell=True, capture_output=True, text=True, timeout=120)
        dur = round(time.time() - t0, 1)
        out = (p.stdout or "").strip()
        clean_lines = [l.strip() for l in out.splitlines() if l.strip() and not l.startswith("DEBUG")]
        last_line = clean_lines[-1] if clean_lines else "Címtár kinyerés befejezve."

        counts = get_db_counts()
        users_cnt = counts.get("users") or counts.get("ad_users") or "-"
        groups_cnt = counts.get("groups") or counts.get("ad_groups") or "-"
        events_cnt = counts.get("audit_events") or counts.get("audit_logs") or counts.get("events") or "-"

        if users_cnt == "-":
            m_u = re.search(r'(\d+)\s*(?:user|felhasználó)', out, re.I)
            if m_u: users_cnt = m_u.group(1)
        if groups_cnt == "-":
            m_g = re.search(r'(\d+)\s*(?:group|csoport)', out, re.I)
            if m_g: groups_cnt = m_g.group(1)

        if p.returncode == 0:
            html = f"""<div style="display:flex; flex-direction:column; gap:5px;">
                <div style="display:flex; justify-content:space-between; align-items:center; border-bottom:1px solid rgba(22,101,52,0.25); padding-bottom:4px;">
                    <span style="font-weight:800; font-size:13px;">✅ AD Adatkinyerés Sikeres ({dur} mp)</span>
                    <span style="font-size:11px; background:rgba(22,101,52,0.15); padding:2px 8px; border-radius:12px; font-weight:700;">{ts}</span>
                </div>
                <div style="display:flex; gap:16px; flex-wrap:wrap; margin-top:3px; font-size:12.5px;">
                    <span>👥 <b>{users_cnt}</b> felhasználó</span>
                    <span>🏢 <b>{groups_cnt}</b> csoport</span>
                    <span>⚡ <b>{events_cnt}</b> audit esemény</span>
                    <span>💾 <b>adaudit</b> adatbázis frissítve</span>
                </div>
                <div style="font-size:11.5px; opacity:0.9; margin-top:2px; font-style:italic;">
                    📋 <b>Zárójelentés:</b> {last_line}
                </div>
            </div>"""
            return JSONResponse(status_code=200, content={"status": "SUCCESS", "message": f"Kinyerés lefutott ({dur} mp).", "html": html})
        else:
            err_msg = last_line
            html = f"""<div style="display:flex; flex-direction:column; gap:4px;">
                <div style="display:flex; justify-content:space-between; align-items:center; border-bottom:1px solid #fecaca; padding-bottom:4px;">
                    <span style="font-weight:800; font-size:13px; color:#b91c1c;">❌ AD Kinyerési Hiba ({dur} mp)</span>
                    <span style="font-size:11px; background:#fee2e2; color:#991b1b; padding:2px 8px; border-radius:12px; font-weight:700;">{ts}</span>
                </div>
                <div style="font-size:12px; color:#b91c1c; margin-top:3px;">⚠️ <b>Hibaüzenet:</b> {err_msg}</div>
            </div>"""
            return JSONResponse(status_code=500, content={"status": "ERROR", "message": err_msg, "html": html})
    except Exception as e:
        return JSONResponse(status_code=500, content={"status": "ERROR", "message": str(e)})

@router.post("/sync-ep-now")
def sync_ep_now():
    t0 = time.time()
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    cmd = "python3 /app/easy_project_sync.py 2>&1 || python3 easy_project_sync.py 2>&1"
    try:
        p = subprocess.run(cmd, shell=True, capture_output=True, text=True, timeout=120)
        dur = round(time.time() - t0, 1)
        out = (p.stdout or "").strip()
        clean_lines = [l.strip() for l in out.splitlines() if l.strip() and not l.startswith("DEBUG")]
        last_line = clean_lines[-1] if clean_lines else "Szinkronizáció befejezve."

        counts = get_db_counts()
        tickets_cnt = counts.get("tickets") or counts.get("issues") or counts.get("helpdesk_tickets") or "-"

        if tickets_cnt == "-":
            m_t = re.search(r'(\d+)\s*(?:issue|ticket|jegy|feladat)', out, re.I)
            if m_t: tickets_cnt = m_t.group(1)

        if p.returncode == 0:
            html = f"""<div style="display:flex; flex-direction:column; gap:5px;">
                <div style="display:flex; justify-content:space-between; align-items:center; border-bottom:1px solid rgba(22,101,52,0.25); padding-bottom:4px;">
                    <span style="font-weight:800; font-size:13px;">✅ Easy Project Szinkronizáció Sikeres ({dur} mp)</span>
                    <span style="font-size:11px; background:rgba(22,101,52,0.15); padding:2px 8px; border-radius:12px; font-weight:700;">{ts}</span>
                </div>
                <div style="display:flex; gap:16px; flex-wrap:wrap; margin-top:3px; font-size:12.5px;">
                    <span>🎫 <b>{tickets_cnt}</b> hibajegy / feladat a rendszerben</span>
                    <span>📁 Projekt: <b>#2403</b></span>
                    <span>💾 <b>adaudit</b> táblák szinkronizálva</span>
                </div>
                <div style="font-size:11.5px; opacity:0.9; margin-top:2px; font-style:italic;">
                    📋 <b>Zárójelentés:</b> {last_line}
                </div>
            </div>"""
            return JSONResponse(status_code=200, content={"status": "SUCCESS", "message": f"Szinkronizáció lefutott ({dur} mp).", "html": html})
        else:
            err_msg = last_line
            html = f"""<div style="display:flex; flex-direction:column; gap:4px;">
                <div style="display:flex; justify-content:space-between; align-items:center; border-bottom:1px solid #fecaca; padding-bottom:4px;">
                    <span style="font-weight:800; font-size:13px; color:#b91c1c;">❌ Easy Project Szinkron Hiba ({dur} mp)</span>
                    <span style="font-size:11px; background:#fee2e2; color:#991b1b; padding:2px 8px; border-radius:12px; font-weight:700;">{ts}</span>
                </div>
                <div style="font-size:12px; color:#b91c1c; margin-top:3px;">⚠️ <b>Hibaüzenet:</b> {err_msg}</div>
            </div>"""
            return JSONResponse(status_code=500, content={"status": "ERROR", "message": err_msg, "html": html})
    except Exception as e:
        return JSONResponse(status_code=500, content={"status": "ERROR", "message": str(e)})

@router.post("/github-push")
async def github_push(req: Request):
    try: body = await req.json()
    except Exception: body = {}

    cfg = get_all_settings()
    repo_url = (body.get("repo_url") or cfg["github"].get("repo_url") or "https://github.com/tspcit/ad-tracker.git").strip()
    token = (body.get("token") or cfg["github"].get("token") or "").strip()
    branch = (body.get("branch") or cfg["github"].get("branch") or "main").strip()
    user_comment = (body.get("comment") or "").strip()

    if not token or token.startswith("•••") or len(token) < 10:
        return JSONResponse(status_code=400, content={
            "status": "ERROR",
            "message": "Hiányzik vagy érvénytelen a GitHub Personal Access Token (PAT)! Kérjük adja meg és mentse el a kártyán a push előtt."
        })

    subprocess.run("git config --global --add safe.directory '*'", shell=True)
    subprocess.run("git config --global user.name 'AD Tracker System'", shell=True)
    subprocess.run("git config --global user.email 'adtracker@tspc.hu'", shell=True)

    curr_branch = subprocess.getoutput("git branch --show-current 2>/dev/null").strip()
    if not curr_branch:
        subprocess.run(f"git checkout -b {branch} 2>/dev/null || true", shell=True)
    elif curr_branch != branch:
        subprocess.run(f"git branch -M {branch} 2>/dev/null || true", shell=True)

    tag_out = subprocess.getoutput("git tag -l 'v*' 2>/dev/null | sort -V | tail -n 1").strip()
    if not tag_out and os.path.exists("app/git_version.txt"):
        try:
            with open("app/git_version.txt") as f: tag_out = f.read().strip()
        except: pass
    if not tag_out: tag_out = "v1.10"

    m = re.match(r'v?(\d+)\.(\d+)', tag_out)
    if m:
        major = int(m.group(1))
        minor = int(m.group(2)) + 1
        new_tag = f"v{major}.{minor:02d}"
    else:
        new_tag = "v1.11"

    ts = time.strftime("%Y-%m-%d %H:%M")
    desc = user_comment if user_comment else f"Release {new_tag} frissítés"
    commit_msg = f"{new_tag}: {desc} ({ts})"

    try:
        os.makedirs("app", exist_ok=True)
        with open("app/git_version.txt", "w", encoding="utf-8") as f:
            f.write(new_tag)
    except: pass

    subprocess.run("git add -A 2>/dev/null", shell=True)
    subprocess.run(f'git commit -m "{commit_msg}" --allow-empty 2>/dev/null', shell=True)
    subprocess.run(f'git tag -a "{new_tag}" -m "{commit_msg}" -f 2>/dev/null', shell=True)

    commit_sha = subprocess.getoutput("git rev-parse --short HEAD 2>/dev/null").strip() or "-"

    clean_repo = re.sub(r'^https?://(?:[^@]+@)?github\.com/', 'github.com/', repo_url).strip().rstrip('/')
    if not clean_repo.endswith('.git'): clean_repo += '.git'
    auth_remote = f"https://x-access-token:{token}@{clean_repo}"

    p = subprocess.run(f'git push "{auth_remote}" HEAD:{branch} --tags 2>&1', shell=True, capture_output=True, text=True, timeout=45)
    
    if p.returncode == 0:
        return JSONResponse(status_code=200, content={
            "status": "SUCCESS",
            "message": f"Új verzió ({new_tag}) sikeresen felpusholva a GitHub ({branch}) ágra!",
            "version": new_tag,
            "name": desc,
            "date": ts,
            "commit": commit_sha
        })
    else:
        err = p.stderr or p.stdout or "Ismeretlen push hiba"
        err_clean = err.replace(token, "***").strip()
        return JSONResponse(status_code=500, content={
            "status": "ERROR",
            "message": f"GitHub Push sikertelen: {err_clean}"
        })

@router.post("/debug-exec")
async def debug_exec(req: Request):
    from auth import COOKIE_NAME, verify_session_token
    token = req.cookies.get(COOKIE_NAME)
    user, _ = verify_session_token(token) if token else (None, None)
    if not user:
        return JSONResponse(status_code=401, content={"status": "ERROR", "message": "Nincs jogosultsága a parancsfuttatáshoz!"})

    try:
        body = await req.json()
    except Exception:
        body = {}
    cmd = (body.get("command") or "").strip().lower()

    if cmd == "health":
        now_ts = time.strftime("%Y-%m-%d %H:%M:%S")
        py_v = subprocess.getoutput("python3 -V")
        out = "Rendszer állapot: OK | Idő: " + now_ts + " | Python: " + py_v
    elif cmd == "docker":
        out = subprocess.getoutput("docker ps 2>/dev/null || ps aux | head -n 10")
    elif cmd == "git":
        out = subprocess.getoutput("git log -n 5 --oneline 2>&1")
    elif cmd == "network":
        out = subprocess.getoutput("ping -c 2 8.8.8.8 2>&1 || curl -I https://google.com 2>&1 | head -n 5")
    else:
        out = "Ismeretlen vagy tiltott diagnosztikai parancs. Engedélyezett: health, docker, git, network"

    return {"status": "SUCCESS", "output": out}

last_runs = {"ad": time.time(), "ep": time.time()}

def background_sync_worker():
    time.sleep(15)
    while True:
        try:
            cfg = get_all_settings()
            ad_min = int(cfg["ad"].get("sync_interval_minutes") or 15)
            ep_min = int(cfg["easyproject"].get("sync_interval_minutes") or 15)
            now = time.time()

            if ad_min > 0 and (now - last_runs["ad"]) >= (ad_min * 60):
                last_runs["ad"] = now
                subprocess.run("python3 /app/sync_engine.py 2>/dev/null || python3 sync_engine.py 2>/dev/null", shell=True)

            if ep_min > 0 and (now - last_runs["ep"]) >= (ep_min * 60):
                last_runs["ep"] = now
                subprocess.run("python3 /app/easy_project_sync.py 2>/dev/null || python3 easy_project_sync.py 2>/dev/null", shell=True)
        except Exception: pass
        time.sleep(45)

# [MIGRÁLVA]: threading háttérszál leállítva